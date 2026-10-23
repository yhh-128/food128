import io
import os
from datetime import datetime, timezone

import pytest
from helpers import SUBS, T, add, png_bytes, run, score
from openpyxl import load_workbook
from PIL import Image

from gift_agent import storage
from gift_agent.storage import (
    ExcelWriteError, GiftStore, ImageError, NotFound, StorageError, visible_rows,
)


@pytest.fixture
def store(cfg):
    return GiftStore(cfg)


def sheet_rows(cfg, name):
    return list(load_workbook(cfg.excel_path)[name].iter_rows(values_only=True))


# ---------------------------------------------------------------- 照片（防護 3）


def test_convert_png_to_real_jpeg():
    assert Image.open(io.BytesIO(storage.convert_to_jpeg(png_bytes()))).format == "JPEG"


@pytest.mark.parametrize("data", [b"", b"hello", b"%PDF-1.4 not an image"])
def test_convert_rejects_non_images(data):
    with pytest.raises(ImageError):
        storage.convert_to_jpeg(data)


def test_photo_name_format_and_collision(tmp_path):
    d = tmp_path / "photos" / "111"
    names = [storage.make_photo_path(d, T, "手織圍巾").name for _ in range(3)]
    assert names == ["2026-10-23 14-05-33 手織圍巾.jpg", "2026-10-23 14-05-33 手織圍巾-2.jpg", "2026-10-23 14-05-33 手織圍巾-3.jpg"]


@pytest.mark.parametrize("name", [
    "../../evil", "..\\..\\evil", "a/b\\c", "a:b*c?d\"e<f>g|h", "CON", "nul", "com1.txt", "", "   ", "..", "...",
    "\x00\x01\n", "name.", "x" * 200, "C:\\Windows\\System32", "/etc/passwd",
])
def test_hostile_names_stay_inside_the_users_folder(tmp_path, name):
    d = tmp_path / "photos" / "111"
    p = storage.make_photo_path(d, T, name)
    assert p.resolve().parent == d.resolve()
    assert all(c not in p.name for c in '<>:"/\\|?*') and ".." not in p.name
    assert len(p.name) < 80


def test_save_photo_goes_into_a_folder_per_user(cfg):
    cfg.inbox_dir.mkdir(parents=True)
    staging = cfg.inbox_dir / "1.jpg"
    staging.write_bytes(b"jpeg")
    rel = storage.save_photo(staging, cfg, "111", T, "手織圍巾")
    assert rel == "photos/111/2026-10-23 14-05-33 手織圍巾.jpg"
    assert (cfg.data_dir / rel).read_bytes() == b"jpeg" and not staging.exists()


@pytest.mark.parametrize("bad_id", ["", "abc", "../1", "1/2", "1 2"])
def test_user_id_must_be_digits_for_folder_names(cfg, bad_id):
    with pytest.raises(StorageError):
        storage.save_photo(cfg.inbox_dir / "x.jpg", cfg, bad_id, T, "x")


# ---------------------------------------------------------------- Excel 文字（防護 4）


@pytest.mark.parametrize("text", ["=1+1", "+1", "-1", "@SUM(A1)", "\t=x", "\r=x", "=HYPERLINK(\"http://x\",\"y\")"])
def test_formula_like_text_becomes_plain_text(text):
    assert storage._prepare_cells([text]) == ["'" + text]


def test_normal_values_untouched_and_illegal_chars_removed():
    assert storage._prepare_cells(["禮物", 5, None, 0.5, "a=b", "送 -1"]) == ["禮物", 5, None, 0.5, "a=b", "送 -1"]
    assert storage._prepare_cells(["a\x00b\x07c"]) == ["abc"]


def test_formula_text_is_text_in_all_three_sheets(store, cfg):
    g = run(store.add_gift(
        user_id="111", uploader="=cmd|' /C calc'!A0", logged_at=T, name="=1+1", category="其他", brand="@x",
        price=1, confidence="low", occasion="+1", recipient="-1", handmade=None, budget=None, score=score(),
        photo="photos/111/a.jpg", said="=HYPERLINK(\"x\")",
    ))
    run(store.add_preference("111", "=bad", "=pref", "=said"))
    run(store.update_gift("111", "=bad", g.number, {"note": "=note"}, None, "=said"))
    for sheet in ("禮物櫃", "修改歷史", "偏好"):
        for row in load_workbook(cfg.excel_path)[sheet].iter_rows(min_row=2):
            assert all(c.data_type != "f" for c in row), sheet


# ---------------------------------------------------------------- 禮物櫃與編號


def test_creates_workbook_with_three_sheets_and_headers(store, cfg):
    run(add(store))
    wb = load_workbook(cfg.excel_path)
    assert wb.sheetnames == ["禮物櫃", "修改歷史", "偏好"]
    assert [c.value for c in wb["禮物櫃"][1]][:6] == ["編號", "送出日期", "登錄時間", "上傳者", "使用者ID", "禮物名稱"]


def test_numbers_are_per_user_and_deleted_numbers_are_not_reused(store):
    a1, a2 = run(add(store, "111")), run(add(store, "111"))
    b1 = run(add(store, "222"))
    assert (a1.number, a2.number, b1.number) == (1, 2, 1)
    run(store.delete_gift("111", "小明", 2))
    assert run(add(store, "111")).number == 3


def test_append_does_not_overwrite_and_stores_user_id_as_text(store, cfg):
    run(add(store, "111", name="第一件"))
    run(add(store, "1234567890123456789", name="第二件"))
    rows = sheet_rows(cfg, "禮物櫃")
    assert [r[5] for r in rows[1:]] == ["第一件", "第二件"]
    assert rows[2][4] == "1234567890123456789"  # 19 位數的 Discord ID 不能被 Excel 變成科學記號


def test_row_contains_scores_and_context(store, cfg):
    g = run(add(store))
    assert g.score == 7.8
    row = sheet_rows(cfg, "禮物櫃")[1]
    assert row[14] == 7.8 and row[15:19] == (6, 9, 8, 8) and row[11] == "媽媽" and row[12] == "是" and row[21] == "有效"


# ---------------------------------------------------------------- 個人隔離（防護 5）


def test_users_only_see_their_own_gifts(store):
    run(add(store, "111", name="A 的圍巾"))
    run(add(store, "222", name="B 的馬克杯"))
    assert [g.name for g in run(store.list_gifts("111"))] == ["A 的圍巾"]
    assert [g.name for g in run(store.list_gifts("222"))] == ["B 的馬克杯"]
    assert run(store.list_gifts("333")) == []


def test_other_users_numbers_do_not_exist_for_me(store):
    run(add(store, "111", name="A 的圍巾"))
    assert run(store.get_gift("222", 1)) is None
    with pytest.raises(NotFound):
        run(store.update_gift("222", "B", 1, {"note": "偷改"}, None, "x"))
    with pytest.raises(NotFound):
        run(store.delete_gift("222", "B", 1))
    with pytest.raises(NotFound):
        run(store.set_reply_id("222", 1, 5))
    assert run(store.get_gift("111", 1)).note is None and run(store.get_gift("111", 1)).status == "有效"


def test_reply_lookup_only_works_for_the_owner(store):
    g = run(add(store, "111"))
    run(store.set_reply_id("111", g.number, 987654321))
    assert run(store.find_gift_by_reply("111", 987654321)).number == 1
    assert run(store.find_gift_by_reply("222", 987654321)) is None


def test_preferences_are_private_and_numbered_per_user(store):
    run(store.add_preference("111", "A", "喜歡實用型"))
    run(store.add_preference("222", "B", "喜歡花"))
    run(store.add_preference("111", "A", "不喜歡甜食"))
    assert [(p.number, p.text) for p in run(store.list_preferences("111"))] == [(1, "喜歡實用型"), (2, "不喜歡甜食")]
    assert [(p.number, p.text) for p in run(store.list_preferences("222"))] == [(1, "喜歡花")]
    with pytest.raises(NotFound):
        run(store.delete_preference("222", "B", 2))  # A 的第 2 條，對 B 來說不存在
    run(store.delete_preference("111", "A", 1))
    assert [p.text for p in run(store.list_preferences("111"))] == ["不喜歡甜食"]


def test_visible_rows_is_the_only_filter_and_handles_deleted():
    class R:
        def __init__(self, uid, status="有效"):
            self.user_id, self.status = uid, status

    rows = [R("1"), R("2"), R("1", "已刪除")]
    assert len(visible_rows(rows, 1)) == 1
    assert len(visible_rows(rows, "1", include_deleted=True)) == 2
    assert visible_rows(rows, "3") == []


# ---------------------------------------------------------------- 修改、刪除、歷史


def test_update_writes_old_and_new_values_to_history(store, cfg):
    g = run(add(store, "111"))
    old, new = run(store.update_gift("111", "小明", g.number, {"recipient": "外婆", "handmade": False}, None, "這是送外婆的"))
    assert (old.recipient, new.recipient, new.handmade) == ("媽媽", "外婆", "否")
    hist = sheet_rows(cfg, "修改歷史")
    assert [h[4] for h in hist[1:]] == ["新增", "修改", "修改"]
    assert hist[2][5:9] == ("送給誰", "媽媽", "外婆", "這是送外婆的")
    assert hist[3][5:8] == ("是否手作", "是", "否")


def test_rescore_updates_scores_and_logs_it(store, cfg):
    g = run(add(store, "111"))
    new_scores = score({"value": 9, "effort": 9, "fit": 9, "uniqueness": 9}, "重新評分後")
    old, new = run(store.update_gift("111", "小明", g.number, {}, new_scores, "重新評分"))
    assert (old.score, new.score, new.s_value, new.reasons) == (7.8, 9.0, 9, "重新評分後")
    last = sheet_rows(cfg, "修改歷史")[-1]
    assert last[4] == "重新評分" and "7.8" in last[6] and "9.0" in last[7]


def test_unchanged_values_do_not_create_history_noise(store, cfg):
    g = run(add(store, "111"))
    run(store.update_gift("111", "小明", g.number, {"recipient": "媽媽"}, None, "沒變"))
    assert len(sheet_rows(cfg, "修改歷史")) == 2  # 表頭＋新增


def test_delete_is_soft_and_keeps_row_photo_and_history(store, cfg):
    g = run(add(store, "111"))
    run(store.delete_gift("111", "小明", g.number, "刪掉"))
    assert run(store.list_gifts("111")) == [] and run(store.get_gift("111", 1)) is None
    assert run(store.get_gift("111", 1, include_deleted=True)).status == "已刪除"
    row = sheet_rows(cfg, "禮物櫃")[1]
    assert row[21] == "已刪除" and row[22] == "photos/111/a.jpg"
    assert sheet_rows(cfg, "修改歷史")[-1][4:8] == ("刪除", "狀態", "有效", "已刪除")
    with pytest.raises(NotFound):
        run(store.delete_gift("111", "小明", 1))  # 已經刪了


def test_history_is_append_only(store, cfg):
    g = run(add(store, "111"))
    run(store.update_gift("111", "小明", g.number, {"note": "第一次"}, None, ""))
    run(store.update_gift("111", "小明", g.number, {"note": "第二次"}, None, ""))
    notes = [h[7] for h in sheet_rows(cfg, "修改歷史")[1:] if h[5] == "備註"]
    assert notes == ["第一次", "第二次"]


# ---------------------------------------------------------------- 失敗時的行為


def test_retries_when_excel_is_open_then_succeeds(store, cfg, monkeypatch, caplog):
    monkeypatch.setattr(storage, "EXCEL_RETRY_SECONDS", 0)
    real, calls = os.replace, {"n": 0}

    def flaky(src, dst):
        calls["n"] += 1
        if calls["n"] <= 2:
            raise PermissionError("locked by Excel")
        return real(src, dst)

    monkeypatch.setattr(storage.os, "replace", flaky)
    assert run(add(store)).number == 1
    assert calls["n"] == 3 and "目前寫不進去" in caplog.text
    assert len(sheet_rows(cfg, "禮物櫃")) == 2  # 沒有重複寫入
    assert not list(cfg.data_dir.glob("*.tmp"))


def test_gives_up_with_clear_error_when_always_locked(store, cfg, monkeypatch):
    monkeypatch.setattr(storage, "EXCEL_RETRY_SECONDS", 0)
    monkeypatch.setattr(storage, "EXCEL_RETRIES", 3)
    monkeypatch.setattr(storage.os, "replace", lambda s, d: (_ for _ in ()).throw(PermissionError("locked")))
    with pytest.raises(ExcelWriteError):
        run(add(store))
    assert not list(cfg.data_dir.glob("*.tmp"))


def test_corrupt_excel_file_raises_write_error(store, cfg):
    cfg.data_dir.mkdir(parents=True)
    cfg.excel_path.write_bytes(b"this is not a zip file")
    with pytest.raises(ExcelWriteError):
        run(add(store))


def test_reading_before_any_file_exists_does_not_create_one(store, cfg):
    assert run(store.list_gifts("111")) == [] and run(store.list_preferences("111")) == []
    assert not cfg.excel_path.exists()
