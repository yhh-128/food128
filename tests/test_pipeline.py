import asyncio

import pytest
from helpers import T, gift_analysis, png_bytes, run
from openpyxl import load_workbook

from gift_agent import report, storage
from gift_agent.analyzer import AnalysisError, GiftAnalysis
from gift_agent.pipeline import JobQueue, PhotoJob, process_upload
from gift_agent.storage import ExcelWriteError, GiftStore


@pytest.fixture
def store(cfg):
    return GiftStore(cfg)


def job(user="111", data=None, caption="", index=1, count=1, aid=1, uploader="小明"):
    return PhotoJob(data if data is not None else png_bytes(), user, uploader, caption, T, aid, index, count)


def fake(result):
    async def analyze(staging, caption, cfg):
        assert staging.exists()
        if isinstance(result, Exception):
            raise result
        return result
    return analyze


def rows(cfg, sheet="禮物櫃"):
    return list(load_workbook(cfg.excel_path)[sheet].iter_rows(min_row=2, values_only=True))


def test_upload_end_to_end(cfg, store):
    out = run(process_upload(job(), cfg, store, fake(gift_analysis())))
    assert out.ok and out.bind_gift == 1
    for part in ("已登錄到你的禮物櫃：#1 手織圍巾", "手作品", "NT$600", "場合：生日", "送給：媽媽", "預算：NT$500",
                 "心意指數：7.8 / 10", "價值感 6", "共 1 件，平均心意指數 7.8", "回覆這則訊息就可以修改"):
        assert part in out.reply
    row = rows(cfg)[0]
    assert row[3:6] == ("小明", "111", "手織圍巾") and row[14] == 7.8
    assert (cfg.data_dir / row[22]).exists() and row[22] == "photos/111/2026-10-23 14-05-33 手織圍巾.jpg"
    assert [h[4] for h in rows(cfg, "修改歷史")] == ["新增"]
    assert not list(cfg.inbox_dir.iterdir())  # 暫存檔已清掉


def test_total_score_comes_from_the_program_not_the_ai(cfg, store):
    run(process_upload(job(), cfg, store, fake(gift_analysis(sub_scores={"value": 10, "effort": 10, "fit": 10, "uniqueness": 10}))))
    assert rows(cfg)[0][14] == 10.0


@pytest.mark.parametrize("status,expected", [("not_a_gift", "看不出是禮物"), ("unreadable", "太暗或太糊")])
def test_not_a_gift_or_unreadable_is_not_recorded(cfg, store, status, expected):
    out = run(process_upload(job(), cfg, store, fake(GiftAnalysis(status=status, what_is_it="自拍照"))))
    assert not out.ok and expected in out.reply and "自拍照" in out.reply
    assert not cfg.excel_path.exists() and not cfg.photos_dir.exists()


def test_analysis_failure_replies_and_writes_nothing(cfg, store):
    out = run(process_upload(job(), cfg, store, fake(AnalysisError("bad json"))))
    assert not out.ok and out.reply == "分析失敗，請再傳一次"
    assert not cfg.excel_path.exists() and not cfg.photos_dir.exists()


def test_unreadable_image_fails_gracefully(cfg, store):
    out = run(process_upload(job(data=b"not an image"), cfg, store, fake(gift_analysis())))
    assert not out.ok and out.reply == report.FAIL_TEXT


def test_excel_failure_has_its_own_message(cfg, store, monkeypatch):
    async def boom(*a, **k):
        raise ExcelWriteError("locked")

    monkeypatch.setattr(store, "add_gift", boom)
    out = run(process_upload(job(), cfg, store, fake(gift_analysis())))
    assert not out.ok and out.reply == report.EXCEL_FAIL_TEXT


def test_each_user_has_their_own_cabinet_numbers_stats_and_photo_folder(cfg, store):
    run(process_upload(job("111", aid=1), cfg, store, fake(gift_analysis())))
    run(process_upload(job("111", aid=2), cfg, store, fake(gift_analysis(sub_scores={"value": 2, "effort": 2, "fit": 2, "uniqueness": 2}))))
    out = run(process_upload(job("222", aid=3, uploader="小華"), cfg, store, fake(gift_analysis(item_name="馬克杯"))))
    assert "#1 馬克杯" in out.reply and "共 1 件，平均心意指數 7.8" in out.reply  # 不受 111 的兩件影響
    assert sorted(p.parent.name for p in cfg.photos_dir.rglob("*.jpg")) == ["111", "111", "222"]
    assert [r[0] for r in rows(cfg) if r[4] == "111"] == [1, 2]


def test_multiple_photos_same_second_get_different_names(cfg, store):
    run(process_upload(job(aid=1, index=1, count=2), cfg, store, fake(gift_analysis())))
    out = run(process_upload(job(aid=2, index=2, count=2), cfg, store, fake(gift_analysis())))
    assert "第 2 張／共 2 張" in out.reply
    paths = [r[22] for r in rows(cfg)]
    assert len(set(paths)) == 2 and paths[1].endswith("-2.jpg")


def test_caption_reaches_the_analyzer_as_data(cfg, store):
    seen = {}

    async def analyze(staging, caption, cfg_):
        seen["caption"] = caption
        return gift_analysis()

    run(process_upload(job(caption="送給媽媽，手作"), cfg, store, analyze))
    assert seen["caption"] == "送給媽媽，手作"


def test_hostile_gift_name_cannot_escape_or_become_a_formula(cfg, store):
    evil = gift_analysis(item_name="=HYPERLINK(\"x\")/../../..\\evil", brand="@x", recipient="-1")
    out = run(process_upload(job(), cfg, store, fake(evil)))
    assert out.ok
    saved = list(cfg.data_dir.rglob("*.jpg"))
    assert len(saved) == 1 and saved[0].parent == cfg.photos_dir / "111"
    row = rows(cfg)[0]
    assert row[5].startswith("'=") and row[7].startswith("'@") and row[11].startswith("'-")


def test_reply_does_not_leave_mentions_active(cfg, store):
    a = gift_analysis(item_name="@everyone **粗體**", reasons="<@123456789012345678>")
    out = run(process_upload(job(), cfg, store, fake(a)))
    assert "@everyone" not in out.reply and "<@123456789012345678>" not in out.reply


# ---------------------------------------------------------------- 佇列


def test_queue_runs_one_at_a_time_in_order_and_survives_errors():
    active, peak, done = 0, 0, []

    async def handler(item):
        nonlocal active, peak
        active += 1
        peak = max(peak, active)
        await asyncio.sleep(0.01)
        active -= 1
        if item == 2:
            raise RuntimeError("boom")
        done.append(item)

    async def main():
        q = JobQueue(handler)
        q.start()
        for i in range(1, 6):
            q.put(i)
        await q.join()

    asyncio.run(main())
    assert peak == 1 and done == [1, 3, 4, 5]
